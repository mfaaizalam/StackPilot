"""SDD tests: validator, sanitizer, fallback builder, renderer, API flow. Runs on SQLite with LLM_FAKE=1."""
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.core.database import Base, SessionLocal, engine
from app.core.llm import LLMError
from app.core.llm_fake import sample_requirements
from app.models.board import Board
from app.models.chat import Chat
from app.models.column import BoardColumn
from app.models.otp import OTP
from app.models.sdd import SDDDocument, SDDJob
from app.models.srs import Diagram, DiagramJob, SRSDocument, SRSJob
from app.models.task import Task
from app.models.user import User
from app.routes import diagrams as diagram_routes
from app.routes import sdd as sdd_routes
from app.routes import srs as srs_routes
from app.routes.auth import get_current_user
from app.schemas.sdd import MAX_COMPONENTS, DesignModel
from app.services import sdd_fallbacks as F
from app.services import sdd_service
from app.services import sdd_validator as V
from app.services.sdd_renderer import component_diagram, render_sdd_markdown

REQ = sample_requirements()
DESC = "A web system for a university library to manage books, members, loans and fines."


@pytest.fixture(autouse=True)
def fresh_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield


def codes(issues, level=None):
    return {i.code for i in issues if level is None or i.level == level}


def design() -> DesignModel:
    return F.build_design(REQ, {})


# ============================================================ fallback builder
def test_fallback_design_is_valid_and_fully_traceable():
    d = design()
    assert V.validate_design(d, REQ) == []
    allocated = {fr for c in d.components for fr in c.fr_ids}
    assert allocated == {f.id for f in REQ.functional_requirements}
    assert {q.nfr_id for q in d.quality_tactics} == {n.id for n in REQ.non_functional_requirements}
    assert {t.entity for t in d.tables} == {e.name for e in REQ.entities}


def test_fallback_layers_and_patterns_follow_standard_rules():
    d = design()
    by_layer = {}
    for c in d.components:
        by_layer.setdefault(c.layer, []).append(c)
    assert by_layer["presentation"][0].pattern == "MVC"
    assert by_layer["api"][0].pattern == "Facade"
    assert all(c.pattern == "Service Layer" for c in by_layer["service"])
    assert by_layer["data"][0].pattern == "Repository"
    assert all(c.pattern == "Adapter" for c in by_layer["integration"])
    assert len(d.components) <= MAX_COMPONENTS


def test_fallback_foreign_keys_follow_entity_relations():
    loan = next(t for t in design().tables if t.entity == "Loan")
    cols = " ".join(loan.columns)
    assert "member_id: INTEGER NOT NULL REFERENCES members(id)" in cols
    assert "book_id: INTEGER NOT NULL REFERENCES books(id)" in cols
    assert any("idx_loans_member_id" in i for i in loan.indexes)


def test_fallback_endpoints_are_unique_and_rest_shaped():
    d = design()
    keys = [(e.method, e.path) for e in d.endpoints]
    assert len(keys) == len(set(keys))
    assert all(e.path.startswith("/api/v1/") for e in d.endpoints)
    register = next(e for e in d.endpoints if "FR-001" in e.fr_ids)
    assert register.method == "POST" and not register.auth_required


def test_fallback_uses_deployment_technology_and_never_invents_any():
    ctx = {"deployment": [{"name": "App Server", "kind": "server", "technology": "Node.js on Linux", "artifacts": []}]}
    assert F.build_design(REQ, ctx).technology_stack[0].technology == "Node.js on Linux"
    assert all(t.technology == "To be confirmed" for t in design().technology_stack)


# ============================================================ validator
def test_unallocated_fr_is_an_error_and_is_not_silently_patched():
    d = design()
    for c in d.components:
        c.fr_ids = [x for x in c.fr_ids if x != "FR-003"]
    assert "FR_NOT_ALLOCATED" in codes(V.validate_design(d, REQ), "error")
    clean, _ = V.sanitize_design(d, REQ)
    assert "FR_NOT_ALLOCATED" in codes(V.validate_design(clean, REQ), "error")  # still reported


def test_missing_tactic_is_an_error():
    d = design()
    d.quality_tactics = d.quality_tactics[1:]
    assert "NFR_NO_TACTIC" in codes(V.validate_design(d, REQ), "error")


def test_dangling_references_are_errors_then_removed_by_sanitize():
    d = design()
    d.components[2].fr_ids.append("FR-999")
    d.components[2].depends_on.append("CMP-099")
    d.endpoints[0].component_id = "CMP-099"
    found = codes(V.validate_design(d, REQ), "error")
    assert {"UNKNOWN_FR", "UNKNOWN_COMPONENT", "ENDPOINT_UNKNOWN_COMPONENT"} <= found
    clean, warns = V.sanitize_design(d, REQ)
    assert not V.has_errors(V.validate_design(clean, REQ)) or codes(V.validate_design(clean, REQ), "error") <= {"FR_NOT_ALLOCATED"}
    assert "FR-999" not in clean.components[2].fr_ids and warns


def test_dependency_cycle_detected_and_broken():
    d = design()
    svc = [c for c in d.components if c.layer == "service"]
    data = next(c for c in d.components if c.layer == "data")
    data.depends_on.append(svc[0].id)  # service -> data -> service
    assert "COMPONENT_CYCLE" in codes(V.validate_design(d, REQ), "error")
    clean, warns = V.sanitize_design(d, REQ)
    assert "COMPONENT_CYCLE" not in codes(V.validate_design(clean, REQ))
    assert any("Circular dependency" in w for w in warns)


def test_layer_violations_are_warnings():
    d = design()
    web = next(c for c in d.components if c.layer == "presentation")
    data = next(c for c in d.components if c.layer == "data")
    web.depends_on.append(data.id)  # presentation skips the service layer
    d.components[-1].layer = "service"  # keep a service layer present
    issues = V.validate_design(d, REQ)
    assert "LAYER_VIOLATION" in codes(issues, "warning")
    data.depends_on.append(next(c.id for c in d.components if c.layer == "api"))  # data calls UP
    assert sum(i.code == "LAYER_VIOLATION" for i in V.validate_design(d, REQ)) >= 2


def test_duplicate_ids_and_no_service_layer():
    d = design()
    d.components[1].id = d.components[0].id
    assert "DUP_COMPONENT_ID" in codes(V.validate_design(d, REQ), "error")
    for c in d.components:
        c.layer = "data"
    assert "NO_SERVICE_LAYER" in codes(V.validate_design(d, REQ), "error")
    assert codes(V.validate_design(DesignModel(**{**design().model_dump(), "components": []}), REQ)) == {"NO_COMPONENTS"}


# ============================================================ sanitizer
def test_sanitize_trims_to_limits_and_renumbers():
    d = design()
    base = d.components[2]
    d.components += [base.model_copy(update={"id": f"CMP-{100 + i}", "name": f"Extra{i}", "fr_ids": [], "depends_on": []}) for i in range(10)]
    clean, warns = V.sanitize_design(d, REQ)
    assert len(clean.components) == MAX_COMPONENTS and warns
    assert [e.id for e in clean.endpoints] == [f"API-{n:03d}" for n in range(1, len(clean.endpoints) + 1)]


def test_sanitize_adds_missing_tables_and_fixes_paths_and_duplicates():
    d = design()
    d.tables = [t for t in d.tables if t.entity != "Loan"]
    d.endpoints.append(d.endpoints[0].model_copy())            # exact duplicate
    d.endpoints[1].path = "api/v1/books"                       # missing leading slash
    clean, warns = V.sanitize_design(d, REQ)
    assert "Loan" in {t.entity for t in clean.tables}
    assert all(e.path.startswith("/") for e in clean.endpoints)
    assert len({(e.method, e.path.lower()) for e in clean.endpoints}) == len(clean.endpoints)
    assert any("Loan" in w for w in warns)


def test_sanitize_never_mutates_its_input():
    d = design()
    d.components[2].fr_ids.append("FR-999")
    before = d.model_dump()
    V.sanitize_design(d, REQ)
    assert d.model_dump() == before


# ============================================================ renderer
def test_markdown_has_all_ieee_1016_sections_and_traceability():
    md = render_sdd_markdown(design(), REQ, [], [], 1, 3)
    for h in ("## 1. Introduction", "## 2. Design Overview", "## 3. Architectural Design", "## 4. Structural Design",
              "## 5. Behavioural Design", "## 6. Deployment Design", "## 7. Data Design", "## 8. Interface Design",
              "## 9. Quality Attribute Design", "## 10. Design Decisions", "## 12. Design Risks",
              "## Appendix A: Requirements-to-Design Traceability"):
        assert h in md
    assert "based on SRS version 3" in md and "NOT ALLOCATED" not in md and "NO TACTIC" not in md
    assert "```mermaid\nflowchart TD" in md


def test_markdown_marks_gaps_and_stale_diagrams_honestly():
    d = design()
    d.components[2].fr_ids = []
    d.quality_tactics = []
    stale = [{"id": 1, "type": "class", "title": "Class diagram", "stale": True, "mermaid": "classDiagram\n  class A"}]
    md = render_sdd_markdown(d, REQ, [], stale)
    assert "NOT ALLOCATED" in md and "NO TACTIC" in md and "out of date" in md and "classDiagram" in md


def test_component_diagram_labels_are_mermaid_safe():
    d = design()
    d.components[0].name = 'Web "Client" [x] (y)'
    out = component_diagram(d)
    assert '"' not in out.split('["')[1].split('"]')[0] and "CMP_001" in out and "-->" in out


# ============================================================ API + job
def make_client(user_id):
    app = FastAPI()
    for r in (srs_routes.router, diagram_routes.router, sdd_routes.router):
        app.include_router(r)
    app.dependency_overrides[get_current_user] = lambda: type("U", (), {"id": user_id})()
    return TestClient(app)


def seed(with_diagrams=True):
    db = SessionLocal()
    db.add_all([User(id=1, name="A", email="a@a.com"), User(id=2, name="B", email="b@b.com")])
    db.add(Board(id=1, name="Library", description="x", owner_id=1))
    db.commit(); db.close()
    c = make_client(1)
    c.post("/boards/1/srs/generate", json={"description": DESC})
    if with_diagrams:
        c.post("/boards/1/srs/diagrams/generate", json={})
    return c


def gen(c):
    r = c.post("/boards/1/sdd/generate", json={})
    assert r.status_code == 202, r.text
    return c.get(f"/boards/1/sdd/jobs/{r.json()['job_id']}").json()


def test_full_sdd_flow_embeds_the_generated_diagrams():
    c = seed()
    st = gen(c)
    assert st["status"] == "done" and st["sdd_id"], st
    out = c.get("/boards/1/sdd").json()
    assert out["version"] == 1 and out["srs_version"] == 1 and out["source"] == "ai" and not out["srs_outdated"]
    assert out["design"]["components"] and out["design"]["endpoints"]
    md = out["markdown"]
    assert "classDiagram" in md and "sequenceDiagram" in md and "deploymentDiagram" not in md
    assert {d["type"] for d in out["diagrams"]} >= {"class", "sequence", "activity", "state", "usecase"}
    assert not any(i["code"] in ("NO_DIAGRAMS", "STALE_DIAGRAMS") for i in out["issues"])


def test_sdd_without_diagrams_still_generates_with_a_warning():
    c = seed(with_diagrams=False)
    assert gen(c)["status"] == "done"
    out = c.get("/boards/1/sdd").json()
    assert "NO_DIAGRAMS" in {i["code"] for i in out["issues"]}
    assert "No class diagram has been generated yet" in out["markdown"]


def test_regenerating_creates_a_new_version_and_versions_endpoint_lists_both():
    c = seed()
    gen(c); gen(c)
    assert [v["version"] for v in c.get("/boards/1/sdd/versions").json()] == [2, 1]
    assert c.get("/boards/1/sdd?version=1").json()["version"] == 1
    assert c.get("/boards/1/sdd?version=9").status_code == 404


def test_srs_edit_marks_sdd_outdated_and_diagrams_stale():
    c = seed()
    gen(c)
    req = c.get("/boards/1/srs/").json()["requirements"]
    req["title"] = "Edited"
    c.patch("/boards/1/srs/requirements", json=req)
    assert c.get("/boards/1/sdd").json()["srs_outdated"] is True
    gen(c)  # new SDD is built on SRS v2, whose diagrams carried over as stale
    out = c.get("/boards/1/sdd").json()
    assert out["srs_version"] == 2 and not out["srs_outdated"]
    assert "STALE_DIAGRAMS" in {i["code"] for i in out["issues"]}


def test_ai_down_uses_rule_based_design(monkeypatch):
    c = seed()
    def boom(*a, **k): raise LLMError("quota")
    monkeypatch.setattr(sdd_service, "llm_json", boom)
    assert gen(c)["status"] == "done"
    out = c.get("/boards/1/sdd").json()
    assert out["source"] == "fallback" and "AI_UNAVAILABLE" in {i["code"] for i in out["issues"]}
    assert not [i for i in out["issues"] if i["level"] == "error"]


def test_one_repair_call_is_made_and_gaps_that_remain_are_reported(monkeypatch):
    c = seed()
    calls = []
    broken = design()
    for comp in broken.components:
        comp.fr_ids = [x for x in comp.fr_ids if x != "FR-003"]
    def fake(prompt, schema, **k):
        calls.append(k.get("system", "")[:20])
        return broken  # the repair returns the same broken design: the gap must stay visible
    monkeypatch.setattr(sdd_service, "llm_json", fake)
    assert gen(c)["status"] == "done"
    assert len(calls) == 2  # draft + exactly ONE repair
    out = c.get("/boards/1/sdd").json()
    assert any(i["code"] == "FR_NOT_ALLOCATED" and "FR-003" in i["message"] for i in out["issues"])
    assert "NOT ALLOCATED" in out["markdown"]


def test_repair_failure_keeps_first_draft(monkeypatch):
    c = seed()
    n = {"i": 0}
    broken = design(); broken.quality_tactics = []
    def fake(prompt, schema, **k):
        n["i"] += 1
        if n["i"] == 1:
            return broken
        raise LLMError("down")
    monkeypatch.setattr(sdd_service, "llm_json", fake)
    assert gen(c)["status"] == "done"


def test_unexpected_error_marks_job_failed_with_safe_message(monkeypatch):
    c = seed()
    monkeypatch.setattr(sdd_service, "build_design", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("secret db password")))
    st = gen(c)
    assert st["status"] == "failed" and "secret" not in (st["error"] or "")


def test_no_srs_returns_404():
    db = SessionLocal()
    db.add(User(id=1, name="A", email="a@a.com")); db.add(Board(id=1, name="B", owner_id=1)); db.commit(); db.close()
    c = make_client(1)
    assert c.post("/boards/1/sdd/generate", json={}).status_code == 404
    assert c.get("/boards/1/sdd").status_code == 404


def test_other_users_get_404_on_every_route():
    c = seed(); gen(c)
    other = make_client(2)
    jid = c.post("/boards/1/sdd/generate", json={}).json()["job_id"]
    for method, url in (("get", "/boards/1/sdd"), ("get", "/boards/1/sdd/versions"), ("get", "/boards/1/sdd/export.md"),
                        ("get", "/boards/1/sdd/export.docx"), ("get", f"/boards/1/sdd/jobs/{jid}"), ("post", "/boards/1/sdd/generate")):
        r = getattr(other, method)(url, **({"json": {}} if method == "post" else {}))
        assert r.status_code == 404, (url, r.status_code)


def test_active_job_is_reused_and_hourly_limit_applies():
    c = seed()
    db = SessionLocal()
    db.add(SDDJob(id="active-1", board_id=1, user_id=1, srs_id=1, status="running", stage="designing")); db.commit(); db.close()
    assert c.post("/boards/1/sdd/generate", json={}).json()["job_id"] == "active-1"
    db = SessionLocal()
    db.query(SDDJob).delete()
    for i in range(5):
        db.add(SDDJob(id=f"old-{i}", board_id=1, user_id=1, srs_id=1, status="done", stage="done"))
    db.commit(); db.close()
    assert c.post("/boards/1/sdd/generate", json={}).status_code == 429


def test_exports_markdown_and_docx():
    c = seed(); gen(c)
    md = c.get("/boards/1/sdd/export.md")
    assert md.status_code == 200 and "SDD_Library_v1.md" in md.headers["content-disposition"]
    assert md.text.startswith("# Software Design Description")
    dx = c.get("/boards/1/sdd/export.docx")
    assert dx.status_code == 200 and dx.content[:2] == b"PK" and "SDD_Library_v1.docx" in dx.headers["content-disposition"]
    from io import BytesIO
    from docx import Document
    text = "\n".join(p.text for p in Document(BytesIO(dx.content)).paragraphs)
    assert "Software Design Description" in text and "Component Specifications" in text


def test_deleting_a_board_removes_its_sdd_versions():
    c = seed(); gen(c)
    db = SessionLocal()
    db.delete(db.query(Board).filter(Board.id == 1).first()); db.commit()
    assert db.query(SDDDocument).count() == 0
    db.close()