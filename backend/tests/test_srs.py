import os
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.core.database import Base, SessionLocal, engine
from app.core.llm_fake import sample_requirements
from app.models.board import Board
from app.models.column import BoardColumn
from app.models.task import Task
from app.models.user import User
from app.models.otp import OTP
from app.models.chat import Chat
from app.models.srs import SRSDocument, SRSJob, Diagram
from app.routes import srs as srs_routes
from app.routes.auth import get_current_user
from app.core.database import get_db
from app.services.srs_service import render_srs_markdown
from app.services.srs_validator import validate_requirements, has_errors


@pytest.fixture(autouse=True)
def fresh_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield


def codes(issues, level=None):
    return {i.code for i in issues if level is None or i.level == level}


# ---------------- validator ----------------
def test_sample_is_valid():
    assert not has_errors(validate_requirements(sample_requirements()))


def test_uncovered_fr_detected():
    r = sample_requirements()
    for u in r.use_cases:
        u.fr_ids = [x for x in u.fr_ids if x != "FR-006"]
    assert "FR_UNCOVERED" in codes(validate_requirements(r), "error")


def test_duplicate_and_bad_ids():
    r = sample_requirements()
    r.functional_requirements[1].id = "FR-001"
    r.functional_requirements[2].id = "REQ-3"
    c = codes(validate_requirements(r), "error")
    assert "DUP_ID" in c and "BAD_ID" in c


def test_unknown_actor_and_unknown_fr():
    r = sample_requirements()
    r.use_cases[0].actor = "Ghost"
    r.use_cases[1].fr_ids.append("FR-099")
    c = codes(validate_requirements(r), "error")
    assert "UNKNOWN_ACTOR" in c and "UNKNOWN_FR" in c


def test_unused_actor():
    r = sample_requirements()
    r.use_cases = [u for u in r.use_cases if u.actor != "Administrator"]
    r.use_cases.append(r.use_cases[0].model_copy(update={"id": "UC-009"}))
    assert "ACTOR_UNUSED" in codes(validate_requirements(r), "error")


def test_warnings_not_errors():
    r = sample_requirements()
    r.functional_requirements[0].statement = "Users can register."
    r.non_functional_requirements[0].metric = "fast"
    iss = validate_requirements(r)
    assert {"NO_SHALL", "NFR_NOT_MEASURABLE"} <= codes(iss, "warning")
    assert not has_errors(iss)


# ---------------- renderer ----------------
def test_markdown_has_all_sections():
    md = render_srs_markdown(sample_requirements(), [], 3)
    for h in ["## 1. Introduction", "## 2. Overall Description", "## 3. Specific Requirements",
              "### 3.2 Functional Requirements", "### 3.3 Non-Functional Requirements",
              "## 4. Use Cases", "## 5. Data Entities", "Appendix A: Traceability Matrix",
              "Appendix B: Open Issues", "Version 3"]:
        assert h in md, h
    assert "NOT COVERED" not in md
    assert "| FR-003 |" in md


def test_markdown_flags_uncovered():
    r = sample_requirements()
    for u in r.use_cases:
        u.fr_ids = [x for x in u.fr_ids if x != "FR-006"]
    assert "NOT COVERED" in render_srs_markdown(r, validate_requirements(r))


# ---------------- API (owner checks + full job) ----------------
def make_client(user_id):
    app = FastAPI()
    app.include_router(srs_routes.router)
    app.dependency_overrides[get_current_user] = lambda: type("U", (), {"id": user_id})()
    return TestClient(app)


def seed():
    db = SessionLocal()
    db.add_all([User(id=1, name="A", email="a@a.com"), User(id=2, name="B", email="b@b.com")])
    db.add(Board(id=1, name="Library", description="x", owner_id=1))
    db.commit(); db.close()


DESC = "A web system for a university library to manage books, members, loans and fines."


def test_full_flow_and_versions():
    seed()
    c = make_client(1)
    q = c.post("/boards/1/srs/questions", json={"description": DESC})
    assert q.status_code == 200 and len(q.json()["questions"]) >= 3

    r = c.post("/boards/1/srs/generate", json={"description": DESC, "answers": [
        {"question": "Platform?", "answer": "Web"}]})
    assert r.status_code == 202
    job_id = r.json()["job_id"]
    # TestClient runs background tasks before returning, so the job is finished
    st = c.get(f"/boards/1/srs/jobs/{job_id}").json()
    assert st["status"] == "done", st
    doc = c.get("/boards/1/srs/").json()
    assert doc["version"] == 1 and "## 1. Introduction" in doc["markdown"]

    # second generation -> version 2
    c.post("/boards/1/srs/generate", json={"description": DESC})
    assert [v["version"] for v in c.get("/boards/1/srs/versions").json()] == [2, 1]
    assert c.get("/boards/1/srs/?version=1").json()["version"] == 1

    # user edit creates a new version and re-validates
    req = doc["requirements"]
    req["title"] = "Edited Title"
    e = c.patch("/boards/1/srs/requirements", json=req)
    assert e.status_code == 200 and e.json()["version"] == 3 and "Edited Title" in e.json()["markdown"]

    md = c.get("/boards/1/srs/export.md")
    assert md.status_code == 200 and "attachment" in md.headers["content-disposition"]


def test_other_user_gets_404_everywhere():
    seed()
    c = make_client(1)
    c.post("/boards/1/srs/generate", json={"description": DESC})
    intruder = make_client(2)
    assert intruder.post("/boards/1/srs/questions", json={"description": DESC}).status_code == 404
    assert intruder.post("/boards/1/srs/generate", json={"description": DESC}).status_code == 404
    assert intruder.get("/boards/1/srs/").status_code == 404
    assert intruder.get("/boards/1/srs/versions").status_code == 404
    assert intruder.get("/boards/1/srs/export.md").status_code == 404
    assert intruder.get("/boards/1/srs/jobs/whatever").status_code == 404


def test_validation_of_input():
    seed()
    c = make_client(1)
    assert c.post("/boards/1/srs/generate", json={"description": "too short"}).status_code == 422


def test_job_failure_is_reported(monkeypatch):
    seed()
    from app.core.llm import LLMError
    from app.services import srs_service
    def boom(*a, **k): raise LLMError("quota exceeded")
    monkeypatch.setattr(srs_service, "build_requirements", boom)
    c = make_client(1)
    job = c.post("/boards/1/srs/generate", json={"description": DESC}).json()["job_id"]
    st = c.get(f"/boards/1/srs/jobs/{job}").json()
    assert st["status"] == "failed" and "quota" in st["error"]


def test_board_delete_cascades_srs():
    seed()
    c = make_client(1)
    c.post("/boards/1/srs/generate", json={"description": DESC})
    db = SessionLocal()
    db.delete(db.query(Board).get(1)); db.commit()
    assert db.query(SRSDocument).count() == 0
    db.close()


def test_docx_export():
    seed()
    c = make_client(1)
    c.post("/boards/1/srs/generate", json={"description": DESC})
    r = c.get("/boards/1/srs/export.docx")
    assert r.status_code == 200 and ".docx" in r.headers["content-disposition"]
    from io import BytesIO
    from docx import Document
    d = Document(BytesIO(r.content))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "1. Introduction" in heads and "4. Use Cases" in heads
    assert len(d.tables) >= 4
    assert make_client(2).get("/boards/1/srs/export.docx").status_code == 404
