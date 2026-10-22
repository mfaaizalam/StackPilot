"""Shared Word layout for the SRS and SDD exports.

Plain, conventional engineering-document look (ISO/IEC/IEEE 29148 and IEEE 1016 style): A4, Times New
Roman body text, Arial headings, a cover page, document control tables, a real table of contents,
numbered captions, bordered tables with a repeating shaded header row, and a running header/footer
with "Page X of Y". Everything is ordinary Word content (styles, tables, fields): fully editable.
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

BODY = "Times New Roman"
HEAD = "Arial"
GREY = "D9D9D9"
LIGHT = "F2F2F2"
PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)
MARGIN = Cm(2.54)


def _set_font(style, name: str, size: float, bold=None, italic=None, color=None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(color or "000000")
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _border(parent, tag: str, sz=4, color="7F7F7F", val="single") -> None:
    el = OxmlElement(tag)
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    parent.append(el)


def _para_border_bottom(p, sz=6, color="000000") -> None:
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    _border(b, "w:bottom", sz, color)
    ppr.append(b)


def _field(paragraph, instr: str, cached: str = "") -> None:
    """A complex field (PAGE, NUMPAGES, SEQ, TOC ...). `cached` is what shows until Word updates it."""
    def run(kind=None, text=None, instr_text=None):
        r = paragraph.add_run()
        if kind:
            c = OxmlElement("w:fldChar")
            c.set(qn("w:fldCharType"), kind)
            r._r.append(c)
        if instr_text:
            i = OxmlElement("w:instrText")
            i.set(qn("xml:space"), "preserve")
            i.text = instr_text
            r._r.append(i)
        if text is not None:
            r.text = text
        return r
    run("begin")
    run(instr_text=f" {instr} ")
    run("separate")
    run(text=cached)
    run("end")


def _ptab_right(paragraph) -> None:
    """Right-aligned tab relative to the margins: stays correct in portrait AND landscape sections."""
    r = paragraph.add_run()
    pt = OxmlElement("w:ptab")
    pt.set(qn("w:relativeTo"), "margin")
    pt.set(qn("w:alignment"), "right")
    pt.set(qn("w:leader"), "none")
    r._r.append(pt)


class Word:
    """Thin helper around python-docx that applies the house style."""

    def __init__(self, title: str, subtitle: str, project: str, version: str, standard: str,
                 doc_code: str, status: str = "Draft", author: str = ""):
        self.d = Document()
        self.title, self.project, self.version, self.standard = title, project, version, standard
        self.doc_code, self.status, self.author = doc_code, status, author
        self.tables = 0
        self.figures = 0
        self._landscape = False
        self._setup_styles()
        self._setup_page(self.d.sections[0])
        self._header_footer(self.d.sections[0])
        self._cover(subtitle)

    # ------------------------------------------------------------ setup
    def _setup_styles(self) -> None:
        st = self.d.styles
        _set_font(st["Normal"], BODY, 11)
        pf = st["Normal"].paragraph_format
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.12
        st["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size, before, after in (("Heading 1", 15, 20, 8), ("Heading 2", 12.5, 14, 6), ("Heading 3", 11, 10, 4)):
            _set_font(st[name], HEAD, size, bold=True, italic=False)
            f = st[name].paragraph_format
            f.space_before, f.space_after, f.keep_with_next = Pt(before), Pt(after), True
            f.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_font(st["Title"], HEAD, 28, bold=True)
        # python-docx's default Title style draws a blue rule: remove it
        ppr = st["Title"].element.get_or_add_pPr()
        for b in ppr.findall(qn("w:pBdr")):
            ppr.remove(b)
        _set_font(st["Caption"], HEAD, 9.5, bold=True, italic=False)
        st["Caption"].paragraph_format.space_before = Pt(4)
        st["Caption"].paragraph_format.space_after = Pt(8)
        st["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for name in ("List Bullet", "List Number"):
            _set_font(st[name], BODY, 11)
            st[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st[name].paragraph_format.space_after = Pt(3)
        for lv in (1, 2, 3):
            try:
                _set_font(st[f"TOC {lv}"], BODY, 11)
            except KeyError:
                pass
        core = self.d.core_properties
        core.title = f"{self.title}: {self.project}"
        core.subject = self.standard
        core.author = self.author or "StackPilot"
        core.created = datetime.utcnow()
        settings = self.d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    def _setup_page(self, sec, landscape: bool = False) -> None:
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = (PAGE_H, PAGE_W) if landscape else (PAGE_W, PAGE_H)
        sec.left_margin = sec.right_margin = MARGIN if not landscape else Cm(2.0)
        sec.top_margin = Cm(2.54) if not landscape else Cm(2.2)
        sec.bottom_margin = Cm(2.3)
        sec.header_distance = sec.footer_distance = Cm(1.25)

    def _header_footer(self, sec, first: bool = True) -> None:
        """Running header and footer. A right tab stop at the text width keeps the right-hand text flush
        in portrait and landscape sections (each section gets its own)."""
        from docx.enum.text import WD_TAB_ALIGNMENT
        sec.different_first_page_header_footer = first
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        width = sec.page_width - sec.left_margin - sec.right_margin

        def put(par, left: str, right: str, border: str, fields: bool = False):
            par.text = ""
            par.style = self.d.styles["Normal"]          # the built-in Header/Footer styles carry their own tab stops
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            def run(t):
                r = par.add_run(t)
                r.font.name, r.font.size = HEAD, Pt(9)
                return r
            run(left)
            run("\t")
            if fields:
                run("Page ")
                _field(par, "PAGE", "1")
                run(" of ")
                _field(par, "NUMPAGES", "1")
                for r in par.runs:
                    r.font.name, r.font.size = HEAD, Pt(9)
            else:
                run(right)
            bd = OxmlElement("w:pBdr")
            _border(bd, border, 6, "7F7F7F")
            par._p.get_or_add_pPr().append(bd)

        put(sec.header.paragraphs[0], self.project, f"{self.title}  |  Version {self.version}", "w:bottom")
        put(sec.footer.paragraphs[0], f"{self.doc_code}  |  {self.status}", "", "w:top", fields=True)

    def _cover(self, subtitle: str) -> None:
        d = self.d
        for _ in range(5):
            d.add_paragraph()
        p = d.add_paragraph(self.title, style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p = d.add_paragraph()
        r = p.add_run(subtitle)
        r.font.name, r.font.size, r.bold = HEAD, Pt(14), False
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _para_border_bottom(p, 12, "000000")
        p.paragraph_format.space_after = Pt(18)
        p = d.add_paragraph()
        r = p.add_run(self.project)
        r.font.name, r.font.size, r.bold = HEAD, Pt(20), True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(60)
        rows = [("Document ID", self.doc_code), ("Version", self.version), ("Status", self.status),
                ("Date", f"{datetime.utcnow():%d %B %Y}"), ("Standard", self.standard)]
        if self.author:
            rows.append(("Prepared by", self.author))
        self._plain_table(rows, [Cm(4.2), Cm(11.8)], shade_first=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        d.add_page_break()

    def _portrait(self) -> None:
        """Return to a portrait section, but only when portrait content actually follows a landscape figure."""
        if self._landscape:
            sec = self.d.add_section(WD_SECTION.NEW_PAGE)
            self._setup_page(sec, landscape=False)
            self._header_footer(sec, first=False)
            self._landscape = False

    # ------------------------------------------------------------ text
    def h1(self, text: str, page_break: bool = False):
        switched = self._landscape      # the section break already starts a new page
        self._portrait()
        if page_break and not switched:
            self.d.add_page_break()
        return self.d.add_heading(text, 1)

    def h2(self, text: str):
        self._portrait()
        return self.d.add_heading(text, 2)

    def h3(self, text: str):
        self._portrait()
        return self.d.add_heading(text, 3)

    def para(self, text: str = "", bold=False, italic=False, align=None, keep=False, size=None):
        self._portrait()
        p = self.d.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def label_para(self, label: str, value: str):
        self._portrait()
        p = self.d.add_paragraph()
        p.paragraph_format.keep_together = True
        p.add_run(f"{label}: ").bold = True
        p.add_run(value or "Not specified.")
        return p

    def bullets(self, items: list[str], empty: str = "None identified."):
        self._portrait()
        if not items:
            self.para(empty, italic=True)
            return
        for i in items:
            self.d.add_paragraph(str(i), style="List Bullet")

    def steps(self, items: list[str], empty: str = "None."):
        self._portrait()
        """Numbered list typed explicitly so every list restarts at 1."""
        if not items:
            self.para(empty, italic=True)
            return
        for n, i in enumerate(items, start=1):
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(f"{n}.\u00A0\u00A0{i}")

    def mono(self, text: str):
        self._portrait()
        p = self.d.add_paragraph()
        r = p.add_run(text)
        r.font.name, r.font.size = "Consolas", Pt(8.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    # ------------------------------------------------------------ tables
    def _caption(self, kind: str, text: str, above: bool):
        n = self.tables if kind == "Table" else self.figures
        p = self.d.add_paragraph(style="Caption")
        p.add_run(f"{kind} ")
        _field(p, f"SEQ {kind} \\* ARABIC", str(n))
        p.add_run(f": {text}")
        if above:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3)
        return p

    def _fmt_table(self, t, widths, header: bool, size=9.5) -> None:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            _border(b, f"w:{side}", 4, "7F7F7F")
        tblPr.append(b)
        m = OxmlElement("w:tblCellMar")
        for side, v in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            m.append(e)
        tblPr.append(m)
        if widths:
            for ci, col in enumerate(t.columns):
                col.width = widths[ci]
            for old in tblPr.findall(qn("w:tblW")):
                tblPr.remove(old)
            tw = OxmlElement("w:tblW")
            tw.set(qn("w:w"), str(int(sum(widths) / 635)))
            tw.set(qn("w:type"), "dxa")
            tblPr.append(tw)
        for ri, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trpr.append(cs)
            if header and ri == 0:
                th = OxmlElement("w:tblHeader")
                trpr.append(th)
            for ci, cell in enumerate(row.cells):
                if widths:
                    cell.width = widths[ci]
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(size)
                        r.font.name = BODY if not (header and ri == 0) else HEAD
                        if header and ri == 0:
                            r.bold = True
                if header and ri == 0:
                    _shade(cell, GREY)

    def avail(self) -> Emu:
        s = self.d.sections[-1]
        return s.page_width - s.left_margin - s.right_margin

    def _widths(self, weights: list[float]) -> list[Emu]:
        total = float(sum(weights))
        a = self.avail()
        return [Emu(int(a * w / total)) for w in weights]

    def table(self, headers: list[str], rows: list[list], weights: list[float] | None = None,
              caption: str | None = None, size: float = 9.5, empty: str = "None."):
        self._portrait()
        if caption:
            self.tables += 1
            self._caption("Table", caption, above=True)
        if not rows:
            self.para(empty, italic=True)
            return None
        weights = weights or [1] * len(headers)
        t = self.d.add_table(rows=1, cols=len(headers))
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = (str(v) if str(v).strip() else "\u2014")
        self._fmt_table(t, self._widths(weights), True, size)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def _plain_table(self, rows: list[tuple], widths, shade_first=False, align=None):
        t = self.d.add_table(rows=0, cols=len(rows[0]))
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
                if i == 0 and shade_first:
                    _shade(cells[i], LIGHT)
                    for run in cells[i].paragraphs[0].runs:
                        run.bold = True
        self._fmt_table(t, widths, False, 10.5)
        return t

    def kv_table(self, title: str, pairs: list[tuple[str, str]], caption: str | None = None, head_fill: str = GREY):
        """One record as a two-column table with a shaded title row (used for requirements, use cases ...)."""
        self._portrait()
        if caption:
            self.tables += 1
            self._caption("Table", caption, above=True)
        t = self.d.add_table(rows=1, cols=2)
        top = t.rows[0].cells
        m = top[0].merge(top[1])
        m.text = title
        for r in m.paragraphs[0].runs:
            r.bold = True
        _shade(m, head_fill)
        for k, v in pairs:
            c = t.add_row().cells
            c[0].text = k
            c[1].text = str(v) if str(v).strip() else "\u2014"
            _shade(c[0], LIGHT)
            for r in c[0].paragraphs[0].runs:
                r.bold = True
        self._fmt_table(t, self._widths([1.1, 4.0]), False, 9.5)
        for row in t.rows[:1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = HEAD
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def list_cell(self, items: list[str]) -> str:
        return "\n".join(f"\u2022 {i}" for i in items) if items else "\u2014"

    # ------------------------------------------------------------ figures
    def figure(self, png: bytes, caption: str, w_px: float, h_px: float, note: str = ""):
        """Insert a diagram. Wide diagrams get their own landscape page so the text stays legible."""
        wide = w_px / max(h_px, 1) > 1.35 and w_px > 620
        if wide:
            if not self._landscape:
                sec = self.d.add_section(WD_SECTION.NEW_PAGE)
                self._setup_page(sec, landscape=True)
                self._header_footer(sec, first=False)
                self._landscape = True
            else:
                self.d.add_page_break()
            max_w, max_h = Cm(24.5), Cm(15.0)
        else:
            self._portrait()
            max_w, max_h = Cm(15.9), Cm(21.0)
        ratio = w_px / h_px
        w = max_w
        if w / ratio > max_h:
            w = int(max_h * ratio)
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(io.BytesIO(png), width=Emu(int(w)))
        self.figures += 1
        self._caption("Figure", caption, above=False)
        if note:
            self.para(note, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)

    # ------------------------------------------------------------ front matter
    def front_matter(self, history: list[list[str]]) -> None:
        self._portrait()
        d = self.d
        p = d.add_paragraph()
        r = p.add_run("Document Control")
        r.font.name, r.font.size, r.bold = HEAD, Pt(15), True
        p.paragraph_format.keep_with_next = True
        self.para("Revision history", bold=True, keep=True)
        self.table(["Version", "Date", "Author", "Description of change"], history, [1, 1.3, 1.8, 4])
        self.para("Review and approval", bold=True, keep=True)
        blank = [["Prepared by", "", "", ""], ["Reviewed by", "", "", ""], ["Approved by", "", "", ""]]
        self.table(["Role", "Name", "Signature", "Date"], blank, [1.4, 2, 2, 1.2])
        d.add_page_break()
        p = d.add_paragraph()
        r = p.add_run("Table of Contents")
        r.font.name, r.font.size, r.bold = HEAD, Pt(15), True
        toc = d.add_paragraph()
        _field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Right-click here and choose Update Field to build the table of contents.")
        p = d.add_paragraph()
        r = p.add_run("List of Figures")
        r.font.name, r.font.size, r.bold = HEAD, Pt(13), True
        p.paragraph_format.space_before = Pt(18)
        lof = d.add_paragraph()
        _field(lof, 'TOC \\h \\z \\c "Figure"', "Right-click here and choose Update Field to build the list of figures.")
        p = d.add_paragraph()
        r = p.add_run("List of Tables")
        r.font.name, r.font.size, r.bold = HEAD, Pt(13), True
        p.paragraph_format.space_before = Pt(18)
        lot = d.add_paragraph()
        _field(lot, 'TOC \\h \\z \\c "Table"', "Right-click here and choose Update Field to build the list of tables.")
        d.add_page_break()

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()


def add_scene(word: Word, scene, caption: str, stale: bool = False, png: bytes | None = None) -> None:
    """Draw a diagram as a figure. If the user edited it in the app, their edited picture is used."""
    from app.services.diagram_drawio import to_model
    from app.services.diagram_render import to_png

    if png:
        from PIL import Image
        w_px, h_px = Image.open(io.BytesIO(png)).size
        word.figure(png, caption, w_px / 2, h_px / 2, "")
        return
    png = to_png(scene, to_model(scene), scale=2.0)
    note = "This diagram was drawn from an earlier version of the requirements and may be out of date." if stale else ""
    word.figure(png, caption, scene.w, scene.h, note)
