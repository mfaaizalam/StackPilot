"""Word (.docx) export, built from the structured Requirements (not from Markdown)."""
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.schemas.srs import Issue, Requirements


def _table(doc, headers: list[str], rows: list[list[str]]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()


def _bullets(doc, items: list[str], style="List Bullet"):
    if not items:
        doc.add_paragraph("None.")
    for i in items:
        doc.add_paragraph(i, style=style)


def build_docx(req: Requirements, issues: list[Issue], version: int = 1) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Software Requirements Specification", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(req.title)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True
    sub.runs[0].font.size = Pt(16)
    meta = doc.add_paragraph(f"Version {version} · {datetime.utcnow():%Y-%m-%d} · Based on ISO/IEC/IEEE 29148")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("1. Introduction", 1)
    doc.add_heading("1.1 Purpose", 2); doc.add_paragraph(req.purpose)
    doc.add_heading("1.2 Scope", 2); doc.add_paragraph(req.scope)
    doc.add_heading("1.3 Definitions, Acronyms and Abbreviations", 2)
    if req.definitions:
        _table(doc, ["Term", "Definition"], [[d.term, d.definition] for d in req.definitions])
    else:
        doc.add_paragraph("None.")

    doc.add_heading("2. Overall Description", 1)
    doc.add_heading("2.1 Product Perspective", 2); doc.add_paragraph(req.product_perspective)
    doc.add_heading("2.2 User Classes and Characteristics", 2)
    _table(doc, ["Actor", "Description"], [[a.name, a.description] for a in req.actors])
    doc.add_heading("2.3 Constraints", 2); _bullets(doc, req.constraints)
    doc.add_heading("2.4 Assumptions and Dependencies", 2); _bullets(doc, req.assumptions)

    doc.add_heading("3. Specific Requirements", 1)
    doc.add_heading("3.1 External Interfaces", 2); _bullets(doc, req.external_interfaces)

    doc.add_heading("3.2 Functional Requirements", 2)
    by_actor: dict[str, list] = {}
    for f in req.functional_requirements:
        by_actor.setdefault(f.actor, []).append(f)
    for n, (actor, frs) in enumerate(by_actor.items(), start=1):
        doc.add_heading(f"3.2.{n} {actor}", 3)
        for f in frs:
            p = doc.add_paragraph()
            r = p.add_run(f"{f.id} · {f.title}"); r.bold = True
            p.add_run(f"   (priority: {f.priority})").italic = True
            doc.add_paragraph(f.statement)
            if f.acceptance_criteria:
                doc.add_paragraph("Acceptance criteria:").runs[0].italic = True
                _bullets(doc, f.acceptance_criteria)

    doc.add_heading("3.3 Non-Functional Requirements", 2)
    by_cat: dict[str, list] = {}
    for r in req.non_functional_requirements:
        by_cat.setdefault(r.category, []).append(r)
    for n, (cat, rs) in enumerate(by_cat.items(), start=1):
        doc.add_heading(f"3.3.{n} {cat.capitalize()}", 3)
        _table(doc, ["ID", "Requirement", "Measurable target"], [[r.id, r.statement, r.metric] for r in rs])

    doc.add_heading("4. Use Cases", 1)
    for u in req.use_cases:
        doc.add_heading(f"{u.id} · {u.name}", 2)
        doc.add_paragraph(f"Primary actor: {u.actor}")
        doc.add_paragraph(u.description)
        doc.add_paragraph(f"Related requirements: {', '.join(u.fr_ids) or '—'}")
        if u.preconditions:
            doc.add_paragraph("Preconditions").runs[0].bold = True
            _bullets(doc, u.preconditions)
        doc.add_paragraph("Main flow").runs[0].bold = True
        _bullets(doc, u.main_flow, "List Number")
        if u.alternate_flows:
            doc.add_paragraph("Alternate flows").runs[0].bold = True
            _bullets(doc, u.alternate_flows)

    doc.add_heading("5. Data Entities", 1)
    for e in req.entities:
        doc.add_heading(e.name, 2)
        doc.add_paragraph(e.description)
        doc.add_paragraph(f"Attributes: {', '.join(e.attributes) or '—'}")
        doc.add_paragraph(f"Relations: {'; '.join(e.relations) or '—'}")

    doc.add_heading("Appendix A: Traceability Matrix", 1)
    _table(doc, ["Requirement", "Title", "Priority", "Covered by"], [
        [f.id, f.title, f.priority, ", ".join(u.id for u in req.use_cases if f.id in u.fr_ids) or "NOT COVERED"]
        for f in req.functional_requirements
    ])
    doc.add_heading("Appendix B: Open Issues", 1); _bullets(doc, req.open_issues)
    if issues:
        doc.add_heading("Appendix C: Quality Check Findings", 1)
        _bullets(doc, [f"[{i.level.upper()}] {i.message}" for i in issues])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
