"""Word (.docx) export of the SDD, built from the structured DesignModel (not from Markdown).

Diagrams are included as titled Mermaid source listings: rendering them to images needs a headless
browser, which the serverless deployment does not have. Use the Markdown export or the app view for pictures."""
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.schemas.sdd import DesignModel
from app.schemas.srs import Issue, Requirements
from app.services.sdd_renderer import component_diagram


def _table(doc, headers: list[str], rows: list[list[str]]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v) if str(v).strip() else "—"
    doc.add_paragraph()


def _bullets(doc, items: list[str]):
    if not items:
        doc.add_paragraph("None.")
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def _field(doc, label: str, value: str):
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value or "—")


def _code(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def _diagrams(doc, items: list[dict], empty: str):
    if not items:
        doc.add_paragraph(empty)
    for x in items:
        doc.add_paragraph().add_run(x["title"] + (" (out of date)" if x.get("stale") else "")).bold = True
        _code(doc, x.get("mermaid") or "(This diagram could not be rendered.)")


def build_docx(d: DesignModel, req: Requirements, issues: list[Issue], diagrams: list[dict],
               version: int = 1, srs_version: int = 1) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    name = {c.id: c.name for c in d.components}
    of = lambda t: [x for x in diagrams if x["type"] == t]

    title = doc.add_heading("Software Design Description", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(req.title)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True
    sub.runs[0].font.size = Pt(16)
    meta = doc.add_paragraph(f"Version {version} · based on SRS version {srs_version} · {datetime.utcnow():%Y-%m-%d} · Based on IEEE 1016")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("1. Introduction", 1)
    doc.add_heading("1.1 Purpose", 2)
    doc.add_paragraph(f"This document describes how the system specified in the SRS (version {srs_version}) is designed.")
    doc.add_heading("1.2 Scope", 2); doc.add_paragraph(req.scope)
    doc.add_heading("1.3 Definitions, Acronyms and Abbreviations", 2)
    if req.definitions:
        _table(doc, ["Term", "Definition"], [[t.term, t.definition] for t in req.definitions])
    else:
        doc.add_paragraph("None.")

    doc.add_heading("2. Design Overview", 1)
    doc.add_heading("2.1 Architectural Style", 2)
    p = doc.add_paragraph(); p.add_run(d.architecture_style).bold = True
    doc.add_paragraph(d.architecture_overview)
    doc.add_heading("2.2 Technology Stack", 2)
    _table(doc, ["Layer", "Technology", "Rationale"], [[t.layer, t.technology, t.rationale] for t in d.technology_stack])
    doc.add_heading("2.3 Design Constraints and Assumptions", 2)
    _bullets(doc, req.constraints + req.assumptions)

    doc.add_heading("3. Architectural Design", 1)
    doc.add_heading("3.1 Component Overview", 2)
    _table(doc, ["ID", "Component", "Layer", "Pattern", "Implements"],
           [[c.id, c.name, c.layer, c.pattern, ", ".join(c.fr_ids)] for c in d.components])
    doc.add_heading("3.2 Component Dependencies", 2)
    _code(doc, component_diagram(d))
    doc.add_heading("3.3 Component Specifications", 2)
    for c in d.components:
        doc.add_heading(f"{c.id} · {c.name}", 3)
        _field(doc, "Layer", c.layer); _field(doc, "Design pattern", c.pattern)
        _field(doc, "Responsibility", c.responsibility); _field(doc, "Implements", ", ".join(c.fr_ids))
        _field(doc, "Data entities", ", ".join(c.entities))
        _field(doc, "Depends on", ", ".join(f"{x} ({name.get(x, x)})" for x in c.depends_on))

    doc.add_heading("4. Structural Design", 1)
    doc.add_heading("4.1 Use Cases and Actors", 2)
    _table(doc, ["Use case", "Name", "Actor", "Requirements"], [[u.id, u.name, u.actor, ", ".join(u.fr_ids)] for u in req.use_cases])
    doc.add_heading("4.2 Class Diagram", 2)
    _diagrams(doc, of("class"), "No class diagram has been generated yet.")

    doc.add_heading("5. Behavioural Design", 1)
    for n, (h, t) in enumerate((("Sequence Diagrams", "sequence"), ("Activity Diagrams", "activity"), ("State Diagrams", "state")), start=1):
        doc.add_heading(f"5.{n} {h}", 2)
        _diagrams(doc, of(t), f"No {t} diagrams have been generated yet.")

    doc.add_heading("6. Deployment Design", 1)
    _diagrams(doc, of("deployment"), "No deployment diagram has been generated yet.")

    doc.add_heading("7. Data Design", 1)
    for t in d.tables:
        doc.add_heading(t.table, 2)
        _field(doc, "Stores", t.entity)
        _table(doc, ["Column", "Definition"], [[c.partition(":")[0].strip(), c.partition(":")[2].strip()] for c in t.columns])
        if t.indexes:
            _bullets(doc, t.indexes)
        if t.notes:
            doc.add_paragraph(t.notes)

    doc.add_heading("8. Interface Design", 1)
    doc.add_heading("8.1 External Interfaces", 2); _bullets(doc, req.external_interfaces)
    doc.add_heading("8.2 REST API", 2)
    _table(doc, ["ID", "Method", "Path", "Actor", "Handled by", "Auth", "Implements"],
           [[e.id, e.method, e.path, e.actor, name.get(e.component_id, e.component_id), "yes" if e.auth_required else "no",
             ", ".join(e.fr_ids)] for e in d.endpoints])
    doc.add_heading("8.3 Endpoint Details", 2)
    for e in d.endpoints:
        doc.add_heading(f"{e.id} · {e.method} {e.path}", 3)
        doc.add_paragraph(e.summary)
        _field(doc, "Request", e.request); _field(doc, "Response", e.response); _field(doc, "Errors", ", ".join(e.errors))

    doc.add_heading("9. Quality Attribute Design", 1)
    tactic = {q.nfr_id: q for q in d.quality_tactics}
    _table(doc, ["Requirement", "Category", "Target", "Design tactic", "Realized by"],
           [[n.id, n.category, n.metric, tactic[n.id].tactic if n.id in tactic else "NO TACTIC",
             ", ".join(name.get(c, c) for c in tactic[n.id].realized_by) if n.id in tactic else ""]
            for n in req.non_functional_requirements])

    doc.add_heading("10. Design Decisions", 1)
    for dd in d.decisions:
        doc.add_heading(f"{dd.id} · {dd.title}", 2)
        _field(doc, "Decision", dd.decision); _field(doc, "Rationale", dd.rationale)
        _field(doc, "Alternatives rejected", ", ".join(dd.alternatives)); _field(doc, "Driven by", ", ".join(dd.nfr_ids))

    doc.add_heading("11. Error Handling and Cross-Cutting Concerns", 1); _bullets(doc, d.error_handling)
    doc.add_heading("12. Design Risks", 1); _bullets(doc, d.design_risks)

    doc.add_heading("Appendix A: Requirements-to-Design Traceability", 1)
    rows = []
    for f in req.functional_requirements:
        comps = [c for c in d.components if f.id in c.fr_ids]
        rows.append([f.id, f.title, ", ".join(c.name for c in comps) or "NOT ALLOCATED",
                     ", ".join(e.id for e in d.endpoints if f.id in e.fr_ids), ", ".join(sorted({x for c in comps for x in c.entities}))])
    _table(doc, ["Requirement", "Title", "Components", "Endpoints", "Data entities"], rows)

    if issues:
        doc.add_heading("Appendix B: Quality Check Findings", 1)
        _bullets(doc, [f"[{i.level.upper()}] {i.message}" for i in sorted(issues, key=lambda i: i.level != "error")])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()